"""Create clearly synthetic fixtures and exercise every parser end-to-end.

Run: python scripts/demo_multiformat.py --output demo_multiformat [--with-dwg]
Never operates on company documents. Refuses to overwrite an existing demo folder.
"""
import argparse
import json
import shutil
from pathlib import Path
from common import load_config, ROOT
from engineering_data import digest
from pipeline import run_multiformat


def main():
    import ezdxf
    import openpyxl
    import pymupdf
    import ifcopenshell
    from docx import Document
    ap = argparse.ArgumentParser()
    ap.add_argument('--output', required=True)
    ap.add_argument('--with-dwg', action='store_true')
    args = ap.parse_args()
    root = Path(args.output).resolve()
    root.mkdir(parents=True, exist_ok=False)
    inp = root / 'input'; inp.mkdir()
    cad = ezdxf.new(); cad.blocks.new('RTU-01'); cad.modelspace().add_blockref('RTU-01', (1, 2, 3))
    cad.modelspace().add_text('SCADA RTU PANEL')
    cad.saveas(inp / 'drawing.dxf')
    if args.with_dwg: shutil.copyfile(ROOT / 'tests/fixtures/sample.dwg', inp / 'sample.dwg')
    for location, qty in [(inp / 'boq.xlsx', 3), (root / 'old_boq.xlsx', 2)]:
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Equipment'
        ws.append(['SYNTHETIC BOQ']); ws.merge_cells('A1:E1')
        ws.append(['Item No.', 'Model', 'Description', 'Qty', 'Unit'])
        ws.append(['001', 'RTU01', 'Remote Terminal Unit', qty, 'ea'])
        ws.append(['002', 'TRAY-X600', 'Cable Tray', 1, 'ea'])
        wb.save(location)
    pdf = pymupdf.open(); page = pdf.new_page(); page.insert_text((40, 40), '1. SCADA requirements\nRTU shall connect to the UPS panel.')
    pdf.save(inp / 'appendix.pdf'); pdf.close()
    word = Document(); word.add_heading('Scope', 1); word.add_heading('SCADA', 2)
    word.add_paragraph('RTU shall retain source evidence.'); word.add_paragraph('Cable Tray should be coordinated.')
    table = word.add_table(rows=1, cols=2); table.cell(0, 0).text = 'UPS'; table.cell(0, 1).text = 'must be supplied'
    word.save(inp / 'sow.docx')
    model = ifcopenshell.file(schema='IFC4'); guid = ifcopenshell.guid.new()
    model.create_entity('IfcCableCarrierSegment', GlobalId=guid, Name='TRAY_X600', PredefinedType='CABLETRAYSEGMENT')
    model.create_entity('IfcDistributionSystem', GlobalId=ifcopenshell.guid.new(), Name='SCADA')
    model.create_entity('IfcSpace', GlobalId=ifcopenshell.guid.new(), Name='Equipment Room')
    model.write(str(inp / 'model.ifc'))
    xml = f'''<exchange><batchtest><clashtests><clashtest><clashresults>
    <clashresult name="Synthetic clash" guid="c1" status="new" distance="-0.1">
    <clashpoint><pos3f x="1" y="2" z="3"/></clashpoint><clashobjects>
    <clashobject><objectattribute><name>GUID</name><value>{guid}</value></objectattribute></clashobject>
    <clashobject><name>RTU-01</name></clashobject></clashobjects></clashresult>
    </clashresults></clashtest></clashtests></batchtest></exchange>'''
    (inp / 'clashes.xml').write_text(xml, encoding='utf-8')
    before = {str(p):digest(p) for p in inp.iterdir()}
    cfg = dict(load_config(), input_dir=str(inp), output_dir=str(root / 'output'))
    summary = run_multiformat(cfg, compare=root / 'old_boq.xlsx', log=lambda level, msg: print(level, msg))
    assert all(digest(p) == value for p, value in before.items()), 'source modified'
    summary['source_hashes_unchanged'] = True
    (root / 'summary.json').write_text(json.dumps(summary, indent=2, ensure_ascii=False), encoding='utf-8')
    print(json.dumps(summary, indent=2, ensure_ascii=False))
    return summary['exit_code']


if __name__ == '__main__':
    raise SystemExit(main())
