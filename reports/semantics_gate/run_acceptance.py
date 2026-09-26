"""Run real parser paths against existing fixtures plus explicit synthetic acceptance cases.

Creates a new output root; never overwrites the historical audit or source documents.
"""
import argparse
import csv
import html
import json
import sqlite3
import sys
from collections import Counter
from pathlib import Path

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'scripts'))
from common import load_config
from engineering_data import digest
from pipeline import run_multiformat

GUID='0AbCdEfGhIjKlMnOpQrStU'


def fixture_files(folder):
    import ezdxf
    import ifcopenshell
    from docx import Document
    folder.mkdir()
    cad=ezdxf.new()
    for name in ('UPS-01','RTU-01','RACK02'):
        cad.blocks.new(name); block=cad.modelspace().add_blockref(name,(1,2,3))
        if name=='UPS-01': block.add_attrib('SPECIFICATION','Voltage: 220 V')
    cad.saveas(folder/'drawing.dxf')
    (folder/'boq.csv').write_text('Item,Model,Specification,Qty\n1,UPS-01,Voltage: 110 V,1\n2,RTU01,,1\n3,RACK01,,1\n4,UNLISTED,,1\n',encoding='utf-8')
    model=ifcopenshell.file(schema='IFC4')
    model.create_entity('IfcCableCarrierSegment',GlobalId=GUID,Name='Tray-01')
    model.create_entity('IfcSpace',GlobalId=ifcopenshell.guid.new(),Name='Equipment Room')
    model.create_entity('IfcDistributionSystem',GlobalId=ifcopenshell.guid.new(),Name='SCADA')
    model.write(str(folder/'model.ifc'))
    word=Document(); word.add_heading('SCADA supply',1)
    table=word.add_table(rows=2,cols=2)
    table.cell(0,0).text='Equipment'; table.cell(0,1).text='Requirement'
    table.cell(1,0).text='UPS'; table.cell(1,1).text='must be supplied'
    word.save(folder/'sow.docx')
    clashes=[]
    for title,value in [('exact',GUID),('different_case',GUID.swapcase())]:
        clashes.append(f'<clashresult name="{title}" guid="{title}" status="new"><clashobjects><clashobject><objectattribute><name>GUID</name><value>{value}</value></objectattribute></clashobject><clashobject><name>RTU-01</name></clashobject></clashobjects></clashresult>')
    (folder/'clash.xml').write_text('<exchange><clashtests><clashtest><clashresults>'+''.join(clashes)+'</clashresults></clashtest></clashtests></exchange>',encoding='utf-8')


def inspect(output):
    path=output/'database/project.db'
    with sqlite3.connect(path.as_uri()+'?mode=ro',uri=True) as db:
        db.row_factory=sqlite3.Row
        refs=[json.loads(r[0]) for r in db.execute('SELECT payload_json FROM cross_reference_results')]
        objects=[json.loads(r[0]) for r in db.execute('SELECT payload_json FROM engineering_objects')]
        requirements=[json.loads(r[0]) for r in db.execute('SELECT payload_json FROM requirements')]
        nulls=[dict(r) for r in db.execute("SELECT guid,name,tag,system,space,level,width,height,length,x,y,z FROM engineering_objects WHERE source_type='IFC'")]
    with (output/'cross_reference/cross_reference.csv').open(encoding='utf-8-sig',newline='') as f:
        csvrows=list(csv.DictReader(f))
    equal=all({k:'' if r.get(k) is None else str(r[k]) for k in c}==c for r,c in zip(refs,csvrows)) and len(refs)==len(csvrows)
    return dict(database=str(path),db_sha256=digest(path),cross_reference=refs,objects=objects,requirements=requirements,
                sql_null_samples=nulls,result_counts=dict(Counter(r['result'] for r in refs)),csv_db_equal=equal)


def main():
    parser=argparse.ArgumentParser(); parser.add_argument('--output',required=True)
    args=parser.parse_args(); root=Path(args.output).resolve(); root.mkdir(parents=True,exist_ok=False)
    fixtures=root/'acceptance_inputs'; fixture_files(fixtures)
    demo=ROOT/'demo_multiformat'
    sources=list((demo/'input').iterdir())+[demo/'old_boq.xlsx']+list(fixtures.iterdir())
    hashes={str(p):digest(p) for p in sources if p.is_file()}
    historical={str(p):digest(p) for p in (demo/'output').rglob('*') if p.is_file()}
    cfg=load_config()
    existing=run_multiformat(dict(cfg,input_dir=str(demo/'input'),output_dir=str(root/'existing_output')),compare=demo/'old_boq.xlsx')
    corner=run_multiformat(dict(cfg,input_dir=str(fixtures),output_dir=str(root/'acceptance_output')))
    a,b=inspect(root/'existing_output'),inspect(root/'acceptance_output')
    refs=b['cross_reference']
    cases={}
    exact=[r for r in refs if r['check_type']=='CLASH_VS_OBJECT' and r['key']==GUID]
    wrong=[r for r in refs if r['check_type']=='CLASH_VS_OBJECT' and r['key']==GUID.swapcase()]
    cases['different_GUID_case_not_match']=bool(wrong) and all(r['result']=='MISSING_B' for r in wrong)
    cases['same_GUID_exact_match']=bool(exact) and all(r['result']=='EXACT_MATCH' for r in exact)
    req=b['requirements'][0]
    cases['DOCX_row_relationship']=req['subject']=='UPS' and req['original_text']=='must be supplied' and len(json.loads(req['row_context_json']))==2
    for cls in ('IfcSpace','IfcDistributionSystem'):
        rows=[r for r in refs if r['check_type']=='IFC_OBJECT_VS_BOQ' and json.loads(r['evidence_a']).get('object_type')==cls]
        cases[cls+'_not_applicable']=bool(rows) and all(r['result']=='NOT_APPLICABLE' for r in rows)
    related=[r for r in refs if r['check_type']=='CAD_KEYWORD_VS_REQUIREMENT' and r['key']=='UPS']
    cases['keyword_related_not_compliance']=bool(related) and all(r['result']=='RELATED' and r['compliance_status']=='INSUFFICIENT_EVIDENCE' for r in related)
    sample=next(r for r in b['sql_null_samples'] if r['guid']==GUID)
    cases['IFC_missing_SQL_NULL']=all(sample[f] is None for f in ('tag','system','space','level','width','height','length','x','y','z'))
    mismatch=[r for r in refs if r['check_type']=='CAD_OBJECT_VS_BOQ' and r['key']=='UPS01']
    cases['explicit_specification_mismatch']=bool(mismatch) and all(r['result']=='MISMATCH' for r in mismatch)
    fuzzy=[r for r in refs if r['key'] in ('RACK01','RACK02')]
    cases['similar_names_uncertain']=bool(fuzzy) and all(r['result']=='UNCERTAIN' for r in fuzzy)
    changes=[p for p,h in hashes.items() if digest(p)!=h]
    report=dict(scope='Synthetic engineering semantic acceptance only; not production approval',existing_summary=existing,acceptance_summary=corner,
                existing=a,acceptance=b,cases=cases,files_checked=len(hashes),hashes_changed=changes,
                historical_output_unchanged=all(digest(p)==h for p,h in historical.items()),
                production_approval='NOT APPROVED')
    (root/'evidence.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    esc=lambda value:html.escape(str(value),quote=True)
    pieces=['<!doctype html><html lang="zh-Hant"><meta charset="utf-8"><title>Engineering Semantics Gate</title><style>body{font:16px/1.6 system-ui;margin:32px auto;padding:0 24px;max-width:1100px;color:#222}table{border-collapse:collapse;width:100%}td,th{border:1px solid #bbb;padding:8px;text-align:left}pre{white-space:pre-wrap;overflow-wrap:anywhere}a{color:#0756a3}</style><h1>Engineering Semantics Gate</h1><p>限定合成資料驗收。Production Approval: NOT APPROVED。此頁只呈現既有 JSON/SQLite 證據，無 LLM 或重新比對。</p><p><a href="evidence.json">完整證據 JSON</a> · <a href="acceptance_output/database/project.db">project.db</a> · <a href="acceptance_output/cross_reference/cross_reference.csv">cross_reference.csv</a></p><table><tr><th>Case</th><th>Outcome</th></tr>']
    pieces += ['<tr><td>'+esc(k)+'</td><td>'+('PASS' if v else 'FAIL')+'</td></tr>' for k,v in cases.items()]
    pieces.append('</table><h2>各筆比對及兩側證據</h2>')
    for row in refs:
        pieces.append('<details><summary>'+esc(row['check_type']+' | '+row['key']+' | '+row['result']+' | '+row['compliance_status'])+'</summary><pre>'+esc(json.dumps(row,ensure_ascii=False,indent=2))+'</pre></details>')
    pieces.append('</html>'); (root/'index.html').write_text('\n'.join(pieces),encoding='utf-8')
    print(json.dumps(dict(cases=cases,existing=existing,acceptance=corner,files_checked=len(hashes),hashes_changed=len(changes)),ensure_ascii=False,indent=2))
    assert all(cases.values()) and not changes and a['csv_db_equal'] and b['csv_db_equal']
    assert report['historical_output_unchanged'] and not existing['exit_code'] and not corner['exit_code']


if __name__=='__main__': main()
