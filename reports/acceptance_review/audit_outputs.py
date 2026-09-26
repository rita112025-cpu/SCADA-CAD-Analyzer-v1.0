"""Read-only acceptance evidence capture for the existing synthetic demo.

Writes only evidence.json next to this script; never reruns the production pipeline.
Counterexamples run in an in-memory database and are explicitly separate from demo data.
"""
import csv
import json
import sqlite3
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'scripts'))
from engineering_data import Store, SCHEMA, digest
from cross_reference import compare_objects, run
from name_normalizer import normalize_name
from common import load_config


def csv_rows(path):
    with path.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def stringify(value):
    return '' if value is None else str(value)


def main():
    import ezdxf
    import ifcopenshell
    import openpyxl
    import pymupdf
    from docx import Document
    import xml.etree.ElementTree as ET
    out = ROOT / 'demo_multiformat/output'
    database = out / 'database/project.db'
    protected = {str(p): digest(p) for p in out.rglob('*') if p.is_file()}
    c = sqlite3.connect(database.as_uri() + '?mode=ro', uri=True)
    c.row_factory = sqlite3.Row
    tables = {name: [json.loads(r[0]) for r in c.execute(f'SELECT payload_json FROM {name}')] for name in SCHEMA}
    for r in tables['project_files']:
        protected[r['source_file']] = digest(r['source_file'])
    source_checks = [{**r, 'current_hash': digest(r['source_file']), 'hash_matches': digest(r['source_file']) == r['file_hash']} for r in tables['project_files']]
    checks = []
    for name, fields in SCHEMA.items():
        expected = [{k: stringify(row.get(k)) for k in fields} for row in tables[name]]
        actual = csv_rows(out / 'database' / (name + '.csv'))
        checks.append(dict(path='database/' + name + '.csv', rows=len(actual), matches=actual == expected))
    for (report,) in c.execute('SELECT DISTINCT report FROM exports'):
        rows = [json.loads(r[0]) for r in c.execute('SELECT payload_json FROM exports WHERE report=?', (report,))]
        fields = list(dict.fromkeys(k for row in rows for k in row))
        expected = [{k: stringify(row.get(k)) for k in fields} for row in rows]
        actual = csv_rows(out / report)
        checks.append(dict(path=report, rows=len(actual), matches=actual == expected))
    objects = tables['engineering_objects']
    inp = ROOT / 'demo_multiformat/input'
    chains = []
    # Source readers are invoked directly; none writes to the source files.
    cad = ezdxf.readfile(inp / 'drawing.dxf')
    block = next(e for e in cad.modelspace() if e.dxftype() == 'INSERT')
    chains.append(dict(kind='CAD', source=dict(file=str(inp/'drawing.dxf'), handle=block.dxf.handle, name=block.dxf.name, position=list(block.dxf.insert)),
                       parsed=next(r for r in csv_rows(out/'cad/blocks.csv') if r['handle'] == block.dxf.handle),
                       normalized=next(o for o in objects if o.get('handle') == block.dxf.handle and o['source_file'].endswith('drawing.dxf'))))
    wb = openpyxl.load_workbook(inp/'boq.xlsx', read_only=True, data_only=True)
    source_boq = list(next(wb.worksheets[0].iter_rows(min_row=3, max_row=3, values_only=True)))
    wb.close()
    chains.append(dict(kind='BOQ', source=dict(file=str(inp/'boq.xlsx'), location='Equipment!row:3', values=source_boq),
                       parsed=csv_rows(out/'excel/boq_items.csv')[0], normalized=next(r for r in tables['boq_items'] if r['source_file'].endswith('input\\boq.xlsx'))))
    model = ifcopenshell.open(str(inp/'model.ifc'))
    for entity in model.by_type('IfcObject'):
        row = next(o for o in objects if o.get('guid') == entity.GlobalId)
        chains.append(dict(kind='IFC', source=dict(file=str(inp/'model.ifc'), location=f'#{entity.id()}', info=entity.get_info(), assignments=len(getattr(entity, 'HasAssignments', ()))), normalized=row))
    doc = Document(inp/'sow.docx')
    chains.append(dict(kind='DOCX', source=dict(file=str(inp/'sow.docx'), paragraphs=[p.text for p in doc.paragraphs], tables=[[[cell.text for cell in row.cells] for row in t.rows] for t in doc.tables]),
                       parsed=csv_rows(out/'docx/requirements.csv'), normalized=tables['requirements']))
    with pymupdf.open(inp/'appendix.pdf') as pdf:
        chains.append(dict(kind='PDF', source=dict(file=str(inp/'appendix.pdf'), page=1, text=pdf[0].get_text()), parsed=csv_rows(out/'pdf/pdf_hits.csv'), normalized=tables['documents']))
    xml = ET.parse(inp/'clashes.xml').getroot()
    chains.append(dict(kind='Navisworks', source=dict(file=str(inp/'clashes.xml'), text=ET.tostring(xml, encoding='unicode')), parsed=csv_rows(out/'navisworks/navis_clashes.csv'), normalized=tables['clashes']))
    cases = []
    for a, b in [('RTU01','RTU-01'), ('UPS-01','UPS PANEL'), ('RACK02','RACK01')]:
        obj = dict(name=a, source_file='synthetic_probe.dxf', source_type='CAD', **normalize_name(a))
        boq = dict(model=b, source_file='synthetic_probe.xlsx', **normalize_name(b))
        cases.append(dict(case=f'{a} vs {b}', results=list(compare_objects([obj], [boq]))))
    guid = '0AbCdEfGhIjKlMnOpQrStU'
    other = guid.swapcase()
    memory = Store(':memory:')
    memory.add('engineering_objects', dict(name='Test', guid=guid, source_type='IFC', source_file='synthetic_probe.ifc', object_type='IfcCableCarrierSegment'))
    memory.add('clashes', dict(object_a=other, source_file='synthetic_probe.xml'))
    run(memory, load_config())
    cases.append(dict(case='Different case-sensitive IFC GUIDs', input_guid=guid, clash_guid=other,
                      expanded_guids=[ifcopenshell.guid.expand(g) for g in (guid, other)],
                      results=[r for r in memory.rows('cross_reference_results') if r['check_type'] == 'CLASH_VS_OBJECT']))
    memory.db.close()
    nulls = [dict(r) for r in c.execute("SELECT name, typeof(system) AS system_storage, system IS NULL AS system_is_null, quote(system) AS system_literal, quote(space) AS space_literal, quote(level) AS level_literal FROM engineering_objects WHERE source_type='IFC'")]
    data = dict(audited_at=datetime.now(timezone.utc).isoformat(), scope='Existing synthetic demo only; not real project acceptance',
                database=str(database), database_sha256=digest(database), table_counts={k:len(v) for k,v in tables.items()},
                source_checks=source_checks, csv_database_checks=checks, evidence_chains=chains,
                actual_cross_reference=tables['cross_reference_results'], result_counts=dict(Counter(r['result'] for r in tables['cross_reference_results'])),
                ifc_missing_storage=nulls, counterexamples=cases,
                existing_summary=json.loads((ROOT/'demo_multiformat/summary.json').read_text(encoding='utf-8')),
                actual_cad_keyword_hits=len(csv_rows(out/'cad/scada_hits.csv')), actual_pdf_keyword_hits=len(csv_rows(out/'pdf/pdf_hits.csv')),
                generated_source_tree_exists=Path('D:/SCADA').exists())
    c.close()
    data['audit_left_sources_and_outputs_unchanged'] = all(digest(p) == h for p,h in protected.items())
    (Path(__file__).parent/'evidence.json').write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
    print(json.dumps({k:data[k] for k in ('table_counts','result_counts','audit_left_sources_and_outputs_unchanged','actual_cad_keyword_hits','actual_pdf_keyword_hits')}, ensure_ascii=True))
    print('CSV checks:',len(checks),'all equal:',all(r['matches'] for r in checks))
    print('Source hashes:',len(source_checks),'all equal:',all(r['hash_matches'] for r in source_checks))


if __name__ == '__main__':
    main()
