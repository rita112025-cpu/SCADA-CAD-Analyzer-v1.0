from analyze_excel import parse, compare
from engineering_data import digest


def test_merged_workbook_and_quantity(tmp_path, data_store, engineering_cfg):
    import openpyxl
    p = tmp_path / 'boq.xlsx'
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(['Engineering BOQ']); ws.merge_cells('A1:D1')
    ws.append(['Item No.', 'Description', 'Model', 'Qty'])
    ws.append(['1', 'RTU equipment', 'RTU-01', '1,200'])
    wb.create_sheet('General').append(['Unstructured', 'preserved'])
    wb.save(p)
    before = digest(p)
    assert parse(p, data_store, engineering_cfg) == 'OK'
    rows = list(data_store.rows('boq_items'))
    assert rows[0]['quantity'] == 1200 and rows[0]['row_number'] == 3
    assert any(r['section'] == 'General' for r in data_store.rows('document_sections'))
    assert digest(p) == before


def test_revision_not_row_number():
    old = [dict(source_file='old', item_no='1', model='RTU01', description='a', quantity=2)]
    new = [dict(source_file='new', item_no='1', model='RTU-01', description='b', quantity=3)]
    result = list(compare(old, new))[0]['result']
    assert 'QUANTITY_CHANGED' in result and 'DESCRIPTION_CHANGED' in result


def test_duplicate_revision_is_uncertain():
    row = dict(source_file='x', item_no='1', model='RTU01')
    assert list(compare([row, row], [row]))[0]['result'] == 'MATCH_UNCERTAIN'


def test_csv_general_and_chinese(tmp_path, data_store, engineering_cfg):
    p = tmp_path / 'boq.csv'
    p.write_text('項次,品名,數量\n1,UPS,2\n', encoding='utf-8')
    parse(p, data_store, engineering_cfg)
    assert list(data_store.rows('boq_items'))[0]['quantity'] == 2


def test_ragged_csv_header_and_identifiers(tmp_path, data_store, engineering_cfg):
    p = tmp_path / 'boq.csv'
    p.write_text('BOQ title\nItem,Model,Qty\n001,RTU01,3\n')
    parse(p, data_store, engineering_cfg)
    row = next(data_store.rows('boq_items'))
    assert row['item_no'] == '001' and row['quantity'] == 3
