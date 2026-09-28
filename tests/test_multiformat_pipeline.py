import csv
import json
import sqlite3
import sys
import time
from pathlib import Path
import pytest
from engineering_data import digest
from pipeline import run_multiformat


def make_inputs(path):
    import ezdxf
    import fitz
    from docx import Document
    import openpyxl
    path.mkdir(parents=True)
    cad = ezdxf.new(); cad.blocks.new('RTU-01'); cad.modelspace().add_blockref('RTU-01', (1, 2))
    cad.saveas(path / 'drawing.dxf')
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(['Item', 'Model', 'Description', 'Qty']); ws.append(['1','RTU01','RTU equipment',2]); wb.save(path / 'boq.xlsx')
    pdf = fitz.open(); page = pdf.new_page(); page.insert_text((30, 30), 'RTU and UPS'); pdf.save(path / 'spec.pdf'); pdf.close()
    doc = Document(); doc.add_heading('SCADA', 1); doc.add_paragraph('RTU shall be supplied.'); doc.save(path / 'sow.docx')
    from test_navisworks_analyzer import XML
    (path / 'clash.xml').write_text(XML)


def test_mixed_pipeline_isolation(tmp_path, engineering_cfg):
    inp, out = tmp_path / 'input', tmp_path / 'out'
    make_inputs(inp)
    (inp / 'broken.pdf').write_text('broken')
    hashes = {p: digest(p) for p in inp.iterdir()}
    cfg = dict(engineering_cfg, input_dir=str(inp), output_dir=str(out))
    result = run_multiformat(cfg)
    assert result['ok'] == 5 and result['failed'] == 1 and result['exit_code'] == 1
    assert all(digest(p) == h for p, h in hashes.items())
    for rel in ['excel/boq_items.csv','pdf/pdf_hits.csv','docx/requirements.csv','navisworks/navis_clashes.csv','cross_reference/cross_reference.csv','cad/scada_hits.csv']:
        assert (out / rel).exists()
    with sqlite3.connect(out / 'database/project.db') as db:
        assert db.execute('SELECT count(*) FROM project_files').fetchone()[0] == 6
        assert db.execute("SELECT count(*) FROM cross_reference_results WHERE result='NORMALIZED_MATCH' AND compliance_status='NOT_EVALUATED'").fetchone()[0] > 0
    assert 'broken.pdf' in (out / 'logs/errors.log').read_text(encoding='utf-8')


def test_missing_dependency_continues(tmp_path, engineering_cfg, monkeypatch):
    inp = tmp_path / 'input'; inp.mkdir()
    (inp / 'model.ifc').write_text('placeholder'); (inp / 'boq.csv').write_text('Item,Qty\n1,2\n')
    monkeypatch.setitem(sys.modules, 'ifcopenshell', None)
    out = tmp_path / 'out'
    result = run_multiformat(dict(engineering_cfg, input_dir=str(inp), output_dir=str(out)))
    assert result['ok'] == 1 and result['skipped'] == 1
    assert 'SKIPPED_DEPENDENCY' in (out / 'database/project_files.csv').read_text(encoding='utf-8-sig')


def test_rerun_replaces_rows_and_empty_reports(tmp_path, engineering_cfg):
    inp = tmp_path / 'input'; inp.mkdir(); p = inp / 'boq.csv'
    p.write_text('Item,Model,Qty\n1,RTU01,2\n')
    out = tmp_path / 'out'; cfg = dict(engineering_cfg, input_dir=str(inp), output_dir=str(out))
    run_multiformat(cfg); p.write_text('Other\nGeneral\n'); run_multiformat(cfg)
    with (out / 'excel/boq_items.csv').open(encoding='utf-8-sig') as f: assert not list(csv.DictReader(f))
    with sqlite3.connect(out / 'database/project.db') as db:
        assert db.execute('SELECT count(*) FROM project_files').fetchone()[0] == 1


def test_gui_analysis_smoke(tmp_path, engineering_cfg, monkeypatch, tk_root):
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    import app
    monkeypatch.setattr(app.App, '_check_env', lambda self: None)
    root = tk_root; root.withdraw()
    for w in root.winfo_children(): w.destroy()
    try:
        a = app.App(root)
        root.geometry('760x620'); root.update()
        assert a.form_canvas.bbox('all')[3] > 620
        a.form_canvas.yview_moveto(1); root.update()
        assert a.form_canvas.yview()[1] == 1
        inp = tmp_path / 'input'; inp.mkdir(); (inp / 'boq.csv').write_text('Item,Model,Qty\n1,RTU01,2\n')
        a.v_in.set(str(inp)); a.v_out.set(str(tmp_path / 'out'))
        assert not a.env_ok
        assert all(str(b['state']) == 'disabled' for b in a.analysis_outputs.values())
        a._start_analysis()
        deadline = time.time() + 20
        while a.summary is None and time.time() < deadline:
            root.update(); time.sleep(.02)
        assert a.summary and a.summary['ok'] == 1
        assert str(a.analysis_outputs['database/project.db']['state']) == 'normal'
    finally:
        root.deiconify()


def test_source_output_overlap_rejected(tmp_path, engineering_cfg):
    with pytest.raises(ValueError):
        run_multiformat(dict(engineering_cfg, input_dir=str(tmp_path), output_dir=str(tmp_path)))


def test_revision_comparison_excludes_old_from_cross_reference(tmp_path, engineering_cfg):
    inp = tmp_path / 'input'; inp.mkdir()
    (inp / 'new.csv').write_text('Item,Model,Qty\n1,RTU01,3\n')
    old = tmp_path / 'old.csv'; old.write_text('Item,Model,Qty\n1,RTU01,2\n')
    import ezdxf
    cad = ezdxf.new(); cad.blocks.new('RTU-01'); cad.modelspace().add_blockref('RTU-01', (0, 0)); cad.saveas(inp / 'a.dxf')
    out = tmp_path / 'out'
    result = run_multiformat(dict(engineering_cfg, input_dir=str(inp), output_dir=str(out)), compare=old)
    assert result['exit_code'] == 0
    with (out / 'excel/boq_compare.csv').open(encoding='utf-8-sig') as f:
        assert next(csv.DictReader(f))['result'] == 'QUANTITY_CHANGED'
    with sqlite3.connect(out / 'database/project.db') as db:
        assert db.execute("SELECT result FROM cross_reference_results WHERE check_type='CAD_OBJECT_VS_BOQ'").fetchone()[0] == 'NORMALIZED_MATCH'


def test_invalid_csv_is_isolated_and_xml_rollback(tmp_path, engineering_cfg):
    inp = tmp_path / 'input'; inp.mkdir()
    (inp / 'bad.csv').write_bytes(b'\xff\xfe\x00\x80')
    from test_navisworks_analyzer import XML
    (inp / 'bad.xml').write_text(XML + '<broken')
    (inp / 'good.csv').write_text('Item,Qty\n1,2\n')
    out = tmp_path / 'out'
    result = run_multiformat(dict(engineering_cfg, input_dir=str(inp), output_dir=str(out)))
    assert result['failed'] == 2 and result['ok'] == 1
    with sqlite3.connect(out / 'database/project.db') as db:
        assert db.execute('SELECT count(*) FROM clashes').fetchone()[0] == 0


def _snapshot(out):
    return {p.relative_to(out).as_posix(): p.read_bytes() for p in out.rglob('*') if p.is_file() and 'logs' not in p.relative_to(out).parts}


def test_export_failure_keeps_previous_publication(tmp_path, engineering_cfg, monkeypatch):
    """Fault injection: a failure mid-export must leave CSVs and project.db as the previous consistent run."""
    import engineering_data
    inp = tmp_path / 'input'; inp.mkdir(); p = inp / 'boq.csv'
    p.write_text('Item,Model,Qty\n1,RTU01,2\n')
    out = tmp_path / 'out'; cfg = dict(engineering_cfg, input_dir=str(inp), output_dir=str(out))
    run_multiformat(cfg)
    before = _snapshot(out)
    p.write_text('Item,Model,Qty\n1,PLC99,7\n2,UPS02,3\n')
    real, calls = engineering_data.Store._csv, []
    def flaky(path, rows, fields):
        calls.append(path)
        if len(calls) == 5: raise OSError('injected export failure')
        return real(path, rows, fields)
    monkeypatch.setattr(engineering_data.Store, '_csv', staticmethod(flaky))
    with pytest.raises(OSError, match='injected'):
        run_multiformat(cfg)
    assert len(calls) >= 5
    assert _snapshot(out) == before
    assert not list((out / 'database').glob('project_*.db'))


@pytest.mark.parametrize('prior', [True, False], ids=['previous-output', 'fresh-output'])
@pytest.mark.parametrize('fail_at', [2, 'project.db'], ids=['2nd-file', 'project.db'])
def test_publish_move_failure_restores_previous_output(tmp_path, engineering_cfg, monkeypatch, prior, fail_at):
    """Fault injection: a failing os.replace while publishing must restore every output file, leave no staging/backup/temp files, and report the error."""
    import os
    inp = tmp_path / 'input'; inp.mkdir(); p = inp / 'boq.csv'
    p.write_text('Item,Model,Qty\n1,RTU01,2\n')
    out = tmp_path / 'out'; cfg = dict(engineering_cfg, input_dir=str(inp), output_dir=str(out))
    if prior: run_multiformat(cfg)
    before = _snapshot(out) if prior else {}
    p.write_text('Item,Model,Qty\n1,PLC99,7\n2,UPS02,3\n')
    real, forward = os.replace, []
    def flaky(src, dst, *a, **k):
        src, dst = Path(src), Path(dst)
        if ('new' in src.parts and any(x.startswith('.publish_') for x in src.parts)) or (dst.name == 'project.db' and src.name.startswith('project_')):
            forward.append(dst)   # publishing a new file; backup and restore moves are not counted
            if dst.name == 'project.db' if fail_at == 'project.db' else len(forward) == fail_at:
                raise PermissionError(13, 'injected publish failure', str(dst))
        return real(src, dst, *a, **k)
    monkeypatch.setattr(os, 'replace', flaky)
    with pytest.raises(PermissionError, match='injected publish failure'):
        run_multiformat(cfg)
    monkeypatch.undo()
    assert len(forward) >= 2
    assert _snapshot(out) == before
    assert not list((out / 'database').glob('project_*.db'))
    assert not [p for p in out.iterdir() if p.name.startswith('.publish_')]
