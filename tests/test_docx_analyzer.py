from analyze_docx import parse


def test_order_hierarchy_tables_verbatim(tmp_path, data_store, engineering_cfg):
    from docx import Document
    p = tmp_path / 'sow.docx'; doc = Document()
    doc.add_heading('Scope', 1); doc.add_heading('SCADA', 2)
    original = 'RTU shall retain original evidence.'
    doc.add_paragraph(original)
    table = doc.add_table(rows=1, cols=2); table.cell(0, 0).text = 'UPS'; table.cell(0, 1).text = 'must be supplied'
    doc.add_paragraph('An ordinary sentence.')
    doc.save(p)
    parse(p, data_store, engineering_cfg)
    sections = list(data_store.rows('document_sections'))
    assert sections[1]['section'] == 'Scope / SCADA'
    assert sections[3]['kind'] == 'table_row'
    reqs = list(data_store.rows('requirements'))
    assert reqs[0]['original_text'] == original and reqs[0]['requirement_text'] == original
    assert reqs[0]['status'] == 'CANDIDATE' and len(reqs) == 2
