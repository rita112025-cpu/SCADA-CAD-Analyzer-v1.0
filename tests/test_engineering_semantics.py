"""Acceptance gate: correctness before retrieval/LLM integration."""
import json
import sqlite3
import pytest
from cross_reference import clash_results, compare_objects, run
from analyze_docx import parse as docx_parse
from analyze_ifc import parse as ifc_parse


GUID='0AbCdEfGhIjKlMnOpQrStU'


@pytest.mark.parametrize('kind',['GUID','UNKNOWN'])
def test_case_different_guid_cannot_match(kind):
    rows=list(clash_results([dict(guid=GUID,name='tray',source_type='IFC')],
        [dict(object_a=GUID.swapcase(),object_a_type=kind)]))
    assert rows[0]['result']=='MISSING_B'
    assert rows[0]['key']==GUID.swapcase()
    assert rows[0]['confidence']==0


def test_exact_guid_and_no_name_fallback():
    obj=dict(guid=GUID,name='TRAY-01',source_type='IFC')
    row=list(clash_results([obj],[dict(object_a=GUID,object_a_type='GUID')]))[0]
    assert row['result']=='EXACT_MATCH' and row['match_basis']=='GUID_EXACT'
    assert row['compliance_status']=='NOT_EVALUATED'
    row=list(clash_results([obj],[dict(object_a='TRAY-01',object_a_type='GUID')]))[0]
    assert row['result']=='MISSING_B'


def test_handle_is_file_scoped():
    obj=dict(handle='A1',source_file='a.dxf',name='UPS')
    row=list(clash_results([obj],[dict(object_a='A1',object_a_type='HANDLE')]))[0]
    assert row['result']=='UNCERTAIN'
    row=list(clash_results([obj],[dict(object_a='A1',object_a_type='HANDLE',object_a_source='a.dxf')]))[0]
    assert row['result']=='EXACT_MATCH'
    row=list(clash_results([obj],[dict(object_a='A1',object_a_type='HANDLE',object_a_source='b.dxf')]))[0]
    assert row['result']=='MISSING_B'


@pytest.mark.parametrize('cls',['IfcSpace','IfcSystem','IfcDistributionSystem','IfcBuildingStorey'])
def test_reference_objects_not_equipment(cls,data_store,engineering_cfg):
    data_store.add('engineering_objects',dict(name='ROOM',source_file='m.ifc',source_type='IFC',object_type=cls))
    run(data_store,engineering_cfg)
    row=next(data_store.rows('cross_reference_results'))
    assert row['result']=='NOT_APPLICABLE' and row['match_basis']=='REFERENCE_ONLY'


def test_requirement_keyword_is_related(data_store,engineering_cfg):
    data_store.add('engineering_objects',dict(name='UPS-01',source_file='a.dxf',source_type='CAD',object_type='INSERT'))
    data_store.add('requirements',dict(source_file='sow.docx',keyword='UPS',requirement_text='must be supplied',subject='UPS'))
    run(data_store,engineering_cfg)
    rows=[r for r in data_store.rows('cross_reference_results') if r['check_type']=='CAD_KEYWORD_VS_REQUIREMENT']
    assert rows[0]['result']=='RELATED' and rows[0]['compliance_status']=='INSUFFICIENT_EVIDENCE'
    assert rows[0]['confidence']<1


def test_exact_normalized_and_specification_conflict():
    item=dict(source_file='boq.xlsx',model='UPS-01',specification='Voltage: 110 V')
    row=list(compare_objects([dict(name='UPS-01',source_type='CAD',specification='Voltage: 220 V')],[item]))[0]
    assert row['result']=='MISMATCH'
    assert json.loads(row['specification_conflicts_json'])[0]['property']=='voltage'
    assert row['compliance_status']=='NOT_EVALUATED'
    row=list(compare_objects([dict(name='UPS-01',source_type='CAD')],[item]))[0]
    assert row['result']=='EXACT_MATCH'
    row=list(compare_objects([dict(name='UPS01',source_type='CAD')],[item]))[0]
    assert row['result']=='NORMALIZED_MATCH' and row['confidence']<1


def test_incomparable_units_not_mismatch():
    row=list(compare_objects([dict(name='UPS01',source_type='CAD',specification='Power: 1 kW')],
        [dict(model='UPS01',specification='Power: 1000 W')]))[0]
    assert row['result']=='EXACT_MATCH'
    assert not json.loads(row['specification_conflicts_json'])


def test_fuzzy_candidate_is_not_simultaneously_missing():
    rows=list(compare_objects([dict(name='RACK02',source_type='CAD')],[dict(model='RACK01')]))
    assert all(r['result']=='UNCERTAIN' for r in rows)
    assert all(r['confidence']<1 for r in rows)


def test_docx_row_relationships_and_headers(tmp_path,data_store,engineering_cfg):
    from docx import Document
    p=tmp_path/'sow.docx'; doc=Document(); t=doc.add_table(rows=3,cols=3)
    for c,value in zip(t.rows[0].cells,['Equipment','Requirement','Notes']): c.text=value
    for c,value in zip(t.rows[1].cells,['UPS','must be supplied','shall be tested']): c.text=value
    t.cell(2,0).text='RTU'; t.cell(2,1).merge(t.cell(2,2)).text='shall retain evidence'
    doc.save(p); docx_parse(p,data_store,engineering_cfg)
    rows=list(data_store.rows('requirements'))
    assert len(rows)==2
    assert rows[0]['subject']=='UPS' and rows[0]['requirement_text']=='must be supplied'
    assert rows[0]['row_number']==2 and rows[0]['table_index']==1
    assert 'UPS' in rows[0]['keyword']
    fragments=json.loads(rows[0]['requirement_fragments_json'])
    assert [r['text'] for r in fragments]==['must be supplied','shall be tested']
    assert json.loads(rows[0]['column_headers_json'])==['Equipment','Requirement','Notes']
    assert json.loads(rows[1]['merged_cells_json'])
    assert '/row:2' in rows[0]['source_location'] and '/cell:' not in rows[0]['source_location']


def test_docx_vertical_merge_preserves_origin(tmp_path,data_store,engineering_cfg):
    from docx import Document
    p=tmp_path/'sow.docx'; doc=Document(); t=doc.add_table(rows=2,cols=2)
    t.cell(0,0).merge(t.cell(1,0)).text='UPS'
    t.cell(0,1).text='must be supplied'; t.cell(1,1).text='shall be tested'
    doc.save(p); docx_parse(p,data_store,engineering_cfg)
    rows=list(data_store.rows('requirements'))
    assert len(rows)==2 and rows[1]['subject']=='UPS'
    origin=json.loads(rows[1]['row_context_json'])[0]
    assert origin['origin_row']==1 and origin['merged']


def test_ifc_nulls_explicit_empty_and_inheritance(tmp_path,data_store,engineering_cfg):
    ifcopenshell=pytest.importorskip('ifcopenshell')  # optional: requirements-ifc.txt
    m=ifcopenshell.file(schema='IFC4')
    m.create_entity('IfcCableCarrierSegment',GlobalId=GUID)
    m.create_entity('IfcSpace',GlobalId=ifcopenshell.guid.new(),Name='')
    m.create_entity('IfcDuctSegment',GlobalId=ifcopenshell.guid.new(),Name='Duct')
    p=tmp_path/'model.ifc'; m.write(str(p)); ifc_parse(p,data_store,engineering_cfg)
    row=data_store.db.execute('SELECT * FROM engineering_objects WHERE guid=?',(GUID,)).fetchone()
    for field in ('name','tag','system','space','level','width','height','length','x','y','z'): assert row[field] is None
    assert data_store.db.execute("SELECT name FROM engineering_objects WHERE object_type='IfcSpace'").fetchone()[0]==''
    duct=next(r for r in data_store.rows('engineering_objects') if r['object_type']=='IfcDuctSegment')
    assert duct['boq_compare_eligible'] is True


def test_additive_schema_migration_preserves_old_rows(tmp_path):
    from engineering_data import Store
    p=tmp_path/'old.db'; c=sqlite3.connect(p)
    c.execute('CREATE TABLE engineering_objects (name TEXT,payload_json TEXT)')
    c.execute('INSERT INTO engineering_objects VALUES (?,?)',('old','{"name":"old"}')); c.commit(); c.close()
    store=Store(p)
    store.add('engineering_objects',dict(name='new',source_type='IFC',tag=None))
    assert [r['name'] for r in store.rows('engineering_objects')]==['old','new']
    assert store.db.execute("SELECT tag FROM engineering_objects WHERE name='new'").fetchone()[0] is None
    store.db.close()


def test_real_parser_specification_mismatch_and_row_requirement(tmp_path,engineering_cfg):
    import ezdxf
    from docx import Document
    from pipeline import run_multiformat
    from engineering_data import digest
    inp=tmp_path/'input'; inp.mkdir()
    cad=ezdxf.new(); cad.blocks.new('UPS-01')
    block=cad.modelspace().add_blockref('UPS-01',(0,0))
    block.add_attrib('SPECIFICATION','Voltage: 220 V')
    cad.saveas(inp/'drawing.dxf')
    (inp/'boq.csv').write_text('Item,Model,Specification,Qty\n1,UPS-01,Voltage: 110 V,1\n')
    word=Document(); table=word.add_table(rows=1,cols=2)
    table.cell(0,0).text='UPS'; table.cell(0,1).text='must be supplied'; word.save(inp/'sow.docx')
    before={p:digest(p) for p in inp.iterdir()}
    out=tmp_path/'output'; result=run_multiformat(dict(engineering_cfg,input_dir=str(inp),output_dir=str(out)))
    assert result['ok']==3 and result['failed']==0
    with sqlite3.connect(out/'database/project.db') as db:
        row=db.execute("SELECT result,compliance_status,specification_conflicts_json FROM cross_reference_results WHERE check_type='CAD_OBJECT_VS_BOQ'").fetchone()
        assert row[:2]==('MISMATCH','NOT_EVALUATED')
        assert json.loads(row[2])[0]['value_a']=='220'
        assert db.execute("SELECT result FROM cross_reference_results WHERE check_type='CAD_KEYWORD_VS_REQUIREMENT'").fetchone()[0]=='RELATED'
        assert db.execute('SELECT subject FROM requirements').fetchone()[0]=='UPS'
    assert all(digest(p)==h for p,h in before.items())
